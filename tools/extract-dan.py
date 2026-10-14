# -*- coding: utf-8 -*-
"""从段位 xlsx 提取课程数据 → JSON（供 C# 内嵌）。"""
import openpyxl, json, re, sys

sys.stdout.reconfigure(encoding='utf-8')

SRC = r'C:\Users\somer\Downloads\4 6 7k段位 個別ACC計算by gyoza_goki & Blaite1016.xlsx'
OUT = r'C:\Users\somer\Documents\deepseek-harness\default-workspace\OsuTransfer\app\dan-courses.json'

wb = openpyxl.load_workbook(SRC, data_only=False)

def parse_if_chain(formula, ref_cell):
    """解析 IF(sel="key","value",...) 链，返回 [(key, value), ...]。value 可能含转义引号。"""
    if not isinstance(formula, str) or not formula.startswith('='):
        return []
    # 提取所有 "key","value" 对（value 支持 \" 转义）
    pairs = re.findall(r'"((?:[^"\\]|\\.)*)"\s*,\s*"((?:[^"\\]|\\.)*)"', formula)
    return [(k.replace('\\"', '"').replace("\\'", "'"), v.replace('\\"', '"').replace("\\'", "'")) for k, v in pairs]

def parse_cum(formula):
    """从個別ACC公式提取累计音符数 [T1,T2,T3,T4]。"""
    if not isinstance(formula, str) or not formula.startswith('='):
        return []
    branches = re.findall(r'IF\([^,]+,\s*(.*?)(?=,IF\(|\)+$)', formula, re.S)
    # 更稳妥：直接按 IF( 切分支
    parts = re.split(r',IF\(', formula)
    cum = []
    for part in parts:
        m = re.search(r'(\d+)\s*\*', part)
        if m:
            cum.append(int(m.group(1)))
    return cum

result = {}

for sheet_name in ('4k', '6k', '7k'):
    ws = wb[sheet_name]
    courses = []
    r = 5
    while r <= ws.max_row:
        title = ws.cell(row=r, column=2).value
        if title and '段位' in str(title):
            course_name = str(title).replace(' 個別ACC', '').replace('個別ACC', '').strip()
            # 布局: title r, selector D(r+1), header r+2, 曲名 r+3, 通過時ACC r+4, 個別ACC r+5, exact r+6
            sel_ref = ws.cell(row=r + 1, column=4)
            song_cell = ws.cell(row=r + 3, column=3)
            song_pairs = parse_if_chain(song_cell.value, song_cell.coordinate)
            levels = []
            for idx, (level_key, _) in enumerate(song_pairs):
                songs = []
                cum = [0, 0, 0, 0]
                for j, col in enumerate((3, 4, 5, 6)):
                    fcell = ws.cell(row=r + 3, column=col)
                    pairs = parse_if_chain(fcell.value, fcell.coordinate)
                    song_name = ''
                    for k, v in pairs:
                        if k == level_key:
                            song_name = v
                            break
                    songs.append(song_name)
                    ecell = ws.cell(row=r + 6, column=col)
                    cs = parse_cum(ecell.value)
                    if len(cs) == len(pairs) and idx < len(cs):
                        cum[j] = cs[idx]
                # 通過時ACC 当前值（若该块当前选择的正是此 level 则可用，否则空）
                accs = [ws.cell(row=r + 4, column=c).value for c in (3, 4, 5, 6)]
                sel_val = sel_ref.value
                accs_out = accs if sel_val == level_key else [None, None, None, None]
                if any(s == '0' or s == 0 for s in songs):
                    continue  # 模板块，跳过
                if not any(songs):
                    continue
                levels.append({
                    'level': level_key,
                    'songs': songs,
                    'cum': cum,
                    'acc': [round(a, 2) if isinstance(a, (int, float)) else None for a in accs_out],
                })
            if levels:
                courses.append({'course': course_name, 'levels': levels})
                print(f'[{sheet_name}] {course_name}: {len(levels)} levels')
            r += 8
        else:
            r += 1
    result[sheet_name] = courses

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump(result, f, ensure_ascii=False, indent=1)
print('saved:', OUT)
print(json.dumps({k: len(v) for k, v in result.items()}, ensure_ascii=False))
